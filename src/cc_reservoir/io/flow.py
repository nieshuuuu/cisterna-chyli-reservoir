# src/cc_reservoir/io/flow.py
import numpy as np
from openpyxl import load_workbook

def read_time_cc_td(xlsx_path, sheet, header_row):
    """Read contiguous (Time, CC, TD) numeric rows from a sheet starting just below
    header_row (1-based). Stops at the first row whose Time cell is not numeric."""
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb[sheet]
    rows = list(ws.iter_rows(min_row=header_row, values_only=True))
    header = [str(c).strip().lower() if c is not None else "" for c in rows[0]]
    ti, ci, di = header.index("time"), header.index("cc"), header.index("td")
    t, cc, td = [], [], []
    for r in rows[1:]:
        if r[ti] is None or not isinstance(r[ti], (int, float)):
            break
        t.append(float(r[ti])); cc.append(float(r[ci])); td.append(float(r[di]))
    return np.array(t), np.array(cc), np.array(td)
